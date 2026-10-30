# -*- coding: utf-8 -*-
"""Regla 4: descargar la plantilla Excel con los últimos datos y volver a cargarla."""
import base64
import io

from odoo import _, fields, models
from odoo.exceptions import UserError

from ..models.produccion import CONCEPTOS

DICC_CONCEPTOS = dict(CONCEPTOS)  # base; en tiempo de ejecución se usan los vigentes


class ProduccionImportar(models.TransientModel):
    _name = "produccion.importar"
    _description = "Importar / exportar la planificación de producción"

    project_id = fields.Many2one("project.project", string="Proyecto", required=True)
    fichero = fields.Binary(string="Fichero Excel")
    nombre_fichero = fields.Char()
    plantilla = fields.Binary(string="Plantilla generada", readonly=True)
    nombre_plantilla = fields.Char(readonly=True)

    # ------------------------------------------------------------------ exportar
    def action_exportar(self):
        """Genera un Excel con la cuadrícula actual (conceptos × meses)."""
        self.ensure_one()
        try:
            import openpyxl
        except ImportError:
            raise UserError(_("Falta la librería openpyxl en el servidor."))
        p = self.project_id
        lineas = p.produccion_linea_ids
        meses = sorted({l.fecha_mes for l in lineas})
        if not meses:
            raise UserError(_("Genera antes la planificación del proyecto."))
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Producción"
        ws.cell(row=1, column=1, value="Proyecto")
        ws.cell(row=1, column=2, value=p.display_name)
        ws.cell(row=2, column=1, value="Tipo de proyecto")
        ws.cell(row=2, column=2, value=dict(p._fields["tipo_produccion"].selection).get(
            p.tipo_produccion, ""))
        ws.cell(row=3, column=1, value="TCV")
        ws.cell(row=3, column=2, value=p.tcv_efectivo)
        fila0 = 5
        ws.cell(row=fila0, column=1, value="Concepto")
        for j, m in enumerate(meses):
            ws.cell(row=fila0, column=2 + j, value=m.strftime("%m/%Y"))
        indice = {(l.fecha_mes, l.concepto): l for l in lineas}
        for i, (clave, etiqueta) in enumerate(
                self.env["produccion.plan.linea"]._conceptos_vigentes()):
            fila = fila0 + 1 + i
            ws.cell(row=fila, column=1, value=etiqueta)
            for j, m in enumerate(meses):
                l = indice.get((m, clave))
                ws.cell(row=fila, column=2 + j, value=l.importe_previsto if l else 0.0)
        buf = io.BytesIO()
        wb.save(buf)
        self.plantilla = base64.b64encode(buf.getvalue())
        self.nombre_plantilla = "produccion_%s.xlsx" % (p.name or "proyecto")[:40]
        return {"type": "ir.actions.act_window", "res_model": self._name,
                "res_id": self.id, "view_mode": "form", "target": "new"}

    # ------------------------------------------------------------------ importar
    def _parsear_excel(self, fichero_b64):
        """Lee la plantilla (1ª hoja) y devuelve {(mes, concepto): importe}.

        Compartido con el asistente de carga inicial: mismo formato de fichero,
        distinta política sobre los meses cerrados.
        """
        inverso = {e.lower(): c for c, e in
                   self.env["produccion.plan.linea"]._conceptos_vigentes()}
        try:
            import openpyxl
        except ImportError:
            raise UserError(_("Falta la librería openpyxl en el servidor."))
        wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(fichero_b64)), data_only=True)
        ws = wb.active
        fila0 = None
        for r in range(1, 30):
            if str(ws.cell(row=r, column=1).value or "").strip().lower() == "concepto":
                fila0 = r
                break
        if not fila0:
            raise UserError(_("No encuentro la fila de cabecera «Concepto» en el fichero."))
        cols = {}
        for c in range(2, ws.max_column + 1):
            v = ws.cell(row=fila0, column=c).value
            if not v:
                continue
            try:
                if hasattr(v, "year"):
                    cols[c] = fields.Date.to_date(v).replace(day=1)  # datetime→date
                else:
                    mm, aa = str(v).strip().split("/")
                    cols[c] = fields.Date.to_date("%s-%02d-01" % (aa, int(mm)))
            except Exception:
                continue
        valores = {}
        for r in range(fila0 + 1, fila0 + 40):
            etiqueta = str(ws.cell(row=r, column=1).value or "").strip().lower()
            clave = inverso.get(etiqueta)
            if not clave:
                continue
            for c, mes in cols.items():
                val = ws.cell(row=r, column=c).value
                if isinstance(val, (int, float)) and not isinstance(val, bool):
                    valores[(mes, clave)] = float(val)
        return valores

    def action_importar(self):
        """Carga el Excel y actualiza los importes previstos (respetando meses cerrados)."""
        self.ensure_one()
        if not self.fichero:
            raise UserError(_("Selecciona un fichero."))
        valores = self._parsear_excel(self.fichero)
        indice = {(l.fecha_mes, l.concepto): l
                  for l in self.project_id.produccion_linea_ids}
        n = saltadas = 0
        for (mes, clave), val in valores.items():
            linea = indice.get((mes, clave))
            if not linea:
                continue
            if linea.mes_cerrado:          # Regla 3
                saltadas += 1
                continue
            linea.with_context(sin_recalculo=True).importe_previsto = val
            n += 1
        self.project_id._recalcular_ingresos()
        aviso = _("Importadas %s celdas.") % n
        if saltadas:
            aviso += _(" %s celdas de meses cerrados se han respetado sin cambios.") % saltadas
        return {"type": "ir.actions.client", "tag": "display_notification",
                "params": {"title": _("Importación completada"), "message": aviso,
                           "type": "success", "next": {"type": "ir.actions.act_window_close"}}}

