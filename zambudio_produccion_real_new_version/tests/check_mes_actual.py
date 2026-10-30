"""Pruebas locales sin Odoo: ejecutan los métodos reales con registros simulados.

    python tests/check_mes_actual.py

No sustituyen a test_mes_actual.py (ORM/BD de Odoo). No conectan con servidores.
Requieren python-dateutil, pytz y openpyxl, habituales en el entorno de Odoo.
"""
import ast
import base64
import io
import sys
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from pathlib import Path
from types import ModuleType, SimpleNamespace
import unittest
from unittest.mock import patch

import openpyxl
import pytz
from dateutil.relativedelta import relativedelta

ROOT = Path(__file__).resolve().parents[1]


class Date:
    today = date(2026, 10, 30)
    context_today = staticmethod(lambda record: Date.today)
    to_string = staticmethod(lambda value: value.isoformat())
    to_date = staticmethod(date.fromisoformat)


class Datetime:
    now = staticmethod(lambda: datetime(2026, 10, 30, 12))
    to_string = staticmethod(lambda value: value.strftime("%Y-%m-%d %H:%M:%S"))


def load_methods(relative, class_name, extra=None):
    """Carga funciones originales, sin importar/inicializar el registro de Odoo."""
    path = ROOT / relative
    tree = ast.parse(path.read_text(encoding="utf-8"))
    cls = next(n for n in tree.body if isinstance(n, ast.ClassDef) and n.name == class_name)
    cls.bases = []
    cls.body = [n for n in cls.body if isinstance(n, ast.FunctionDef)]
    constants = [n for n in tree.body if isinstance(n, ast.Assign)]
    namespace = dict(
        api=SimpleNamespace(model=lambda fn: fn),
        fields=SimpleNamespace(Date=Date, Datetime=Datetime),
        UserError=ValueError, defaultdict=defaultdict, pytz=pytz,
        datetime=datetime, time=time, timedelta=timedelta,
        relativedelta=relativedelta, io=io, base64=base64,
    )
    namespace.update(extra or {})
    exec(compile(ast.Module(body=constants + [cls], type_ignores=[]), str(path), "exec"), namespace)
    return namespace[class_name], namespace


Motor, motor_constants = load_methods("models/produccion_real.py", "ProduccionRealMotor")
Project, _ = load_methods("models/project.py", "ProjectProject", {
    "TIPOS_COMPRA": motor_constants["TIPOS_COMPRA"],
})
Export, _ = load_methods("wizard/exportar_br.py", "ProduccionExportarBR")


class TestMesActualLocal(unittest.TestCase):
    def setUp(self):
        Date.today = date(2026, 10, 30)
        self.motor = Motor()
        self.motor.env = SimpleNamespace(context={"tz": "Europe/Madrid"}, user=SimpleNamespace(tz=False))
        self.sources = {name: [] for name in ("horas", "compras", "gastos", "almacen")}
        self.queries = []
        for name in self.sources:
            def collect(project, start, end, out, source=name):
                self.queries.append((source, start, end))
                for day, concept, amount in self.sources[source]:
                    if start <= day <= end:
                        out[(day.strftime("%Y-%m"), concept)] += amount
            setattr(self.motor, "_acumular_" + name, collect)
        self.motor.facturacion_por_mes = lambda project: {"2026-10": 200.0}
        self.motor._cuenta_analitica = lambda project: False
        self.motor.horas_requieren_validacion = lambda: True
        self.motor.posibles_duplicados_almacen = lambda *args: []
        self.motor.detalle_horas = lambda *args: []
        self.motor.detalle_compras = lambda *args: []
        self.project = Project()
        p = self.project
        p.env = {"produccion.real.motor.new.version": self.motor}
        p.ensure_one = lambda: None
        p.with_company = lambda company: p
        p._sincronizar_tcv = lambda: None
        p._meses_planificacion = lambda: [date(2026, m, 1) for m in (9, 10, 11)]
        p.company_id = 1
        p.name = "TEST OCTUBRE"
        p.tcv_efectivo, p.tcv_fuente, p.produccion_tcv_alerta = 3000.0, "manual", False
        p.tipo_produccion = "recurrente"
        p._fields = {"tipo_produccion": SimpleNamespace(selection=[("recurrente", "Recurrente")])}
        p.produccion_fecha_inicio, p.produccion_fecha_fin = date(2026, 9, 1), date(2026, 11, 30)
        p.currency_id = SimpleNamespace(symbol="EUR")
        p.produccion_real_last_update = datetime(2026, 10, 1)
        p.produccion_real_stale = True
        p.produccion_linea_ids = []
        p.produccion_avance_ids = []
        for month in (9, 10, 11):
            for concept, prev, real in (("ingreso", 1000.0, 0.0), ("horas_internas", -500.0, -20.0)):
                p.produccion_linea_ids.append(SimpleNamespace(
                    fecha_mes=date(2026, month, 1), concepto=concept,
                    importe_previsto=prev, importe_real=real))
            p.produccion_avance_ids.append(SimpleNamespace(
                fecha_mes=date(2026, month, 1), importe_confirmado=800.0,
                estado="confirmado", importe_sugerido_lineal=1000.0))

    @staticmethod
    def cost(data, concept="horas_internas", month="2026-10"):
        return next(f["celdas"][month] for f in data["costes"] if f["concepto"] == concept)

    def test_project_without_dates_using_actual_module_imports(self):
        """Carga todo project.py: no inyecta UserError en sus variables globales."""
        class PlanningError(Exception):
            pass

        odoo = ModuleType("odoo")
        odoo.fields = SimpleNamespace(Date=Date, Datetime=Datetime)
        odoo.models = SimpleNamespace(Model=object)
        exceptions = ModuleType("odoo.exceptions")
        exceptions.UserError = PlanningError
        motor_module = ModuleType("br_regression.produccion_real")
        motor_module.TIPOS_COMPRA = motor_constants["TIPOS_COMPRA"]
        module = ModuleType("br_regression.project")
        module.__package__ = "br_regression"
        path = ROOT / "models/project.py"
        with patch.dict(sys.modules, {
            "odoo": odoo, "odoo.exceptions": exceptions,
            "br_regression.produccion_real": motor_module,
        }):
            exec(compile(path.read_text(encoding="utf-8"), str(path), "exec"), module.__dict__)

        project = module.ProjectProject()
        project.__dict__.update(self.project.__dict__)
        project.with_company = lambda company: project
        project.produccion_fecha_inicio = project.produccion_fecha_fin = False

        def missing_dates():
            raise PlanningError("Indica las fechas de inicio y fin de la planificación")

        project._meses_planificacion = missing_dates
        data = project.get_produccion_br_data_new_version()
        self.assertEqual(data["meses"], [])
        self.assertEqual(data["estado_mes"], {})
        self.assertEqual(self.queries, [])

    def test_current_costs_refresh_without_persisting_or_closing_income(self):
        self.sources["horas"] = [(Date.today, "horas_internas", -125.0)]
        data = self.project.get_produccion_br_data_new_version()
        self.assertEqual(self.cost(data), {"valor": -125.0, "real": -125.0,
                                          "previsto": -500.0, "desviacion": 375.0})
        self.assertEqual(data["avance"]["2026-10"]["valor"], 1000.0)
        self.assertFalse(data["avance"]["2026-10"]["confirmado"])
        self.assertEqual(data["pm"]["2026-10"], 875.0)
        self.assertEqual(data["pct_pm"]["2026-10"], 87.5)
        self.assertEqual(data["wip"]["2026-10"], 1600.0)
        self.assertEqual(data["estado_mes"], {"2026-09": "pasado", "2026-10": "actual", "2026-11": "futuro"})
        self.sources["horas"].append((Date.today, "horas_internas", -75.0))
        self.assertEqual(self.cost(self.project.get_produccion_br_data_new_version())["valor"], -200.0)
        self.sources["horas"].clear()
        self.assertEqual(self.cost(self.project.get_produccion_br_data_new_version())["valor"], 0.0)
        self.assertEqual(self.project.produccion_real_last_update, datetime(2026, 10, 1))
        self.assertTrue(all(l.importe_real == -20.0 for l in self.project.produccion_linea_ids if l.concepto == "horas_internas"))

    def test_all_sources_missing_plan_line_and_future_date(self):
        self.sources["horas"] = [(Date.today, "horas_externas", -100.0)]
        self.sources["compras"] = [(Date.today, "licencias_software", -250.0)]
        self.sources["gastos"] = [(Date.today, "viajes", -30.0)]
        self.sources["almacen"] = [(Date.today, "materiales", -80.0), (Date.today, "materiales", 20.0),
                                    (date(2026, 10, 31), "materiales", -999.0)]
        data = self.project.get_produccion_br_data_new_version()
        self.assertEqual(sum(f["celdas"]["2026-10"]["valor"] for f in data["costes"]), -440.0)
        self.assertEqual(self.cost(data, "licencias_software")["valor"], -250.0)
        self.assertEqual({name for name, _, _ in self.queries}, set(self.sources))
        self.assertTrue(all((start, end) == (date(2026, 10, 1), Date.today) for _, start, end in self.queries))

    def test_past_and_future_unchanged(self):
        data = self.project.get_produccion_br_data_new_version()
        self.assertEqual(self.cost(data, month="2026-09")["valor"], -20.0)
        self.assertEqual(self.cost(data, month="2026-11")["valor"], -500.0)
        self.assertEqual(data["avance"]["2026-09"]["valor"], 800.0)
        self.assertTrue(data["avance"]["2026-09"]["confirmado"])

    def test_other_year_or_no_dates_does_not_query_current_sources(self):
        self.assertEqual(self.project.get_produccion_br_data_new_version(anio=2025)["meses"], [])
        def missing_dates():
            raise ValueError("Sin fechas")
        self.project._meses_planificacion = missing_dates
        self.assertEqual(self.project.get_produccion_br_data_new_version()["meses"], [])
        self.assertEqual(self.queries, [])

    def test_rollover_and_year_boundary(self):
        Date.today = date(2026, 11, 1)
        self.sources["horas"] = [(Date.today, "horas_internas", -12.0)]
        data = self.project.get_produccion_br_data_new_version()
        self.assertEqual(data["estado_mes"]["2026-10"], "pasado")
        self.assertEqual(self.cost(data, month="2026-11")["valor"], -12.0)
        Date.today = date(2027, 1, 1)
        self.project._meses_planificacion = lambda: [date(2026, 12, 1), date(2027, 1, 1), date(2027, 2, 1)]
        data = self.project.get_produccion_br_data_new_version()
        self.assertEqual(list(data["estado_mes"].values()), ["pasado", "actual", "futuro"])

    def test_datetime_month_limits_and_dst(self):
        domain = self.motor._dominio_fechas_datetime("date", date(2026, 10, 1), date(2026, 10, 31))
        self.assertEqual(domain, [("date", ">=", datetime(2026, 9, 30, 22)),
                                  ("date", "<", datetime(2026, 10, 31, 23))])
        for day, hours in ((date(2026, 3, 29), 23), (date(2026, 10, 25), 25)):
            domain = self.motor._dominio_fechas_datetime("date_done", day, day)
            self.assertEqual((domain[1][2] - domain[0][2]).total_seconds() / 3600, hours)

    def test_today_afternoon_inside_stock_domain(self):
        data = self.project.get_produccion_br_data_new_version()
        self.assertEqual(data["dominio_fechas_actual_datetime"], [
            ("date", ">=", "2026-09-30 22:00:00"), ("date", "<", "2026-10-30 23:00:00")])
        after_midnight = datetime(2026, 10, 30, 20)
        domain = self.motor._dominio_fechas_datetime("date", date(2026, 10, 1), Date.today)
        self.assertTrue(domain[0][2] <= after_midnight < domain[1][2])

    def test_stock_queries_use_datetime_limits(self):
        class Model:
            _fields = {"project_id": True}

            def sudo(self):
                return self

            def search(self, domain, **kwargs):
                self.domain = domain
                return []

        class Env(dict):
            context = {"tz": "Europe/Madrid"}
            user = SimpleNamespace(tz=False)

        move, picking = Model(), Model()
        self.motor.env = Env({"stock.move": move, "stock.picking": picking})
        self.project.id = 42
        self.project.company_id = SimpleNamespace(id=1)
        self.motor._movimientos_compra(self.project, False, date(2026, 10, 1), Date.today)
        Motor._acumular_almacen(self.motor, self.project, date(2026, 10, 1), Date.today, defaultdict(float))
        for model, field in ((move, "date"), (picking, "date_done")):
            self.assertIn((field, ">=", datetime(2026, 9, 30, 22)), model.domain)
            self.assertIn((field, "<", datetime(2026, 10, 30, 23)), model.domain)

    def test_excel_values_labels_colors_and_current_income_is_not_real(self):
        self.sources["horas"] = [(Date.today, "horas_internas", -125.0)]
        wizard = Export()
        wizard.project_id, wizard.anio, wizard.incluir_agregados = self.project, False, True
        wizard.env, wizard.id, wizard._name = self.project.env, 1, "produccion.exportar.br"
        wizard.ensure_one = lambda: None
        wizard.write = lambda values: wizard.__dict__.update(values)
        wizard.action_exportar()
        book = openpyxl.load_workbook(io.BytesIO(base64.b64decode(wizard.fichero)))
        sheet = book["Seguimiento economico"]
        column = next(c.column for c in sheet[5] if c.value == "Oct 26 · En curso")
        row = next(row[0].row for row in sheet if row[0].value == "Horas internas")
        self.assertEqual(sheet.cell(row, column).value, -125.0)
        self.assertEqual(sheet.cell(5, column).fill.fgColor.rgb, "002563EB")
        self.assertEqual(sheet.cell(5, column).font.color.rgb, "00FFFFFF")
        self.assertEqual(sheet.cell(row, column).fill.fgColor.rgb, "00DBEAFE")
        self.assertEqual(sheet.cell(row, sheet.max_column).value, -645.0)
        detail = book["Previsto vs Real"]
        column = next(c.column for c in detail[3] if c.value == "Oct 26 · En curso")
        self.assertIsNone(detail.cell(5, column + 1).value)
        self.assertIsNone(detail.cell(5, column + 2).value)
        self.assertEqual(detail.cell(6, column + 1).value, -125.0)


if __name__ == "__main__":
    unittest.main(verbosity=2)
