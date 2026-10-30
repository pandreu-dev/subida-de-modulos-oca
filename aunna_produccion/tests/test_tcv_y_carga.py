# -*- coding: utf-8 -*-
"""TCV de la ficha, reparto automático y puerta de atrás de meses cerrados."""
from datetime import date

from dateutil.relativedelta import relativedelta

from odoo import fields
from odoo.exceptions import UserError, ValidationError
from odoo.tests import TransactionCase, tagged

from ..models.produccion import PERMISO_MESES_CERRADOS


@tagged("post_install", "-at_install")
class TestTcvYCarga(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.hoy = fields.Date.context_today(cls.env.user)
        cls.mes_actual = cls.hoy.replace(day=1)
        # 6 meses: 3 ya cerrados, el actual y 2 futuros
        cls.ini = cls.mes_actual - relativedelta(months=3)
        cls.fin = (cls.mes_actual + relativedelta(months=3)) - relativedelta(days=1)
        cls.tiene_ficha = "tcv_amount" in cls.env["project.project"]._fields

    def _proyecto(self, tipo="recurring", tcv=6000.0, en_ficha=None):
        vals = {
            "name": "TEST TCV %s" % tipo,
            "zambudio_project_type": tipo,
            "date_start": self.ini,
            "date": self.fin,
        }
        usar_ficha = self.tiene_ficha if en_ficha is None else en_ficha
        if usar_ficha and self.tiene_ficha:
            vals["tcv_amount"] = tcv
        else:
            vals["tcv"] = tcv
        p = self.env["project.project"].create(vals)
        p.action_generar_planificacion()
        return p

    def _ingresos(self, p):
        return {l.fecha_mes: l.importe_previsto
                for l in p.produccion_linea_ids if l.concepto == "ingreso"}

    # ------------------------------------------------------------ TCV
    def test_tcv_efectivo_y_fuente(self):
        p = self._proyecto(tcv=6000.0)
        self.assertAlmostEqual(p.tcv_efectivo, 6000.0, 2)
        self.assertEqual(p.tcv_fuente, "ficha" if self.tiene_ficha else "manual")
        # el reparto recurrente usa el TCV efectivo: cuota 6.000/6 = 1.000 en los
        # meses abiertos; los ya cerrados no se tocan (quedan a 0: para eso está
        # la carga inicial)
        ing = self._ingresos(p)
        self.assertEqual(len(ing), 6)
        for m, v in ing.items():
            self.assertAlmostEqual(v, 0.0 if m < self.mes_actual else 1000.0, 2)
        self.assertAlmostEqual(p.produccion_tcv_aplicado, 6000.0, 2)

    def test_ficha_manda_sobre_manual(self):
        if not self.tiene_ficha:
            self.skipTest("esta base no tiene tcv_amount en la ficha")
        p = self._proyecto(tcv=6000.0, en_ficha=True)
        p.tcv = 999.0  # el manual no debe contar si la ficha tiene TCV
        self.assertAlmostEqual(p.tcv_efectivo, 6000.0, 2)
        self.assertEqual(p.tcv_fuente, "ficha")

    def test_cambio_de_tcv_rehace_reparto(self):
        p = self._proyecto(tcv=6000.0)
        campo = "tcv_amount" if self.tiene_ficha else "tcv"
        p.write({campo: 12000.0})
        # meses cerrados: se quedan como estaban (0); abiertos: se rehacen (2.000)
        ing = self._ingresos(p)
        cerrados = [m for m in ing if m < self.mes_actual]
        abiertos = [m for m in ing if m >= self.mes_actual]
        self.assertEqual(len(cerrados), 3)
        for m in cerrados:
            self.assertAlmostEqual(ing[m], 0.0, 2)
        for m in abiertos:
            self.assertAlmostEqual(ing[m], 2000.0, 2)
        self.assertAlmostEqual(p.produccion_tcv_aplicado, 12000.0, 2)

    def test_tcv_ficha_baja_deja_aviso_sin_bloquear(self):
        if not self.tiene_ficha:
            self.skipTest("esta base no tiene tcv_amount en la ficha")
        p = self._proyecto(tipo="time_materials", tcv=6000.0, en_ficha=True)
        # T&M: ingreso manual de 5.000 en un mes futuro
        futuro = self.mes_actual + relativedelta(months=1)
        linea = p.produccion_linea_ids.filtered(
            lambda l: l.concepto == "ingreso" and l.fecha_mes == futuro)
        linea.importe_previsto = 5000.0
        # la ficha baja el TCV por debajo: NO bloquea, deja aviso
        p.write({"tcv_amount": 4000.0})
        self.assertTrue(p.produccion_tcv_alerta)
        self.assertIn("supera el TCV", p.produccion_tcv_alerta)
        # y al subirlo otra vez, el aviso desaparece
        p.write({"tcv_amount": 6000.0})
        self.assertFalse(p.produccion_tcv_alerta)

    def test_sincronizar_detecta_cambio_por_fuera(self):
        """Simula un TCV de ficha calculado (cambia sin pasar por write)."""
        p = self._proyecto(tcv=6000.0)
        campo = "tcv_amount" if self.tiene_ficha else "tcv"
        self.env.cr.execute("UPDATE project_project SET %s = 3000 WHERE id = %%s" % campo,
                            (p.id,))
        p.invalidate_recordset()
        self.assertAlmostEqual(p.tcv_efectivo, 3000.0, 2)
        self.env["project.project"].cron_sincronizar_tcv()
        ing = self._ingresos(p)
        abiertos = [m for m in ing if m >= self.mes_actual]
        # recurrente: 3.000/6 = 500 por mes en los abiertos
        for m in abiertos:
            self.assertAlmostEqual(ing[m], 500.0, 2)
        self.assertAlmostEqual(p.produccion_tcv_aplicado, 3000.0, 2)

    # ------------------------------------------------------------ mes cerrado dinámico
    def test_mes_cerrado_no_se_congela(self):
        p = self._proyecto()
        l_pasado = p.produccion_linea_ids.filtered(
            lambda l: l.concepto == "ingreso" and l.fecha_mes == self.ini)
        l_actual = p.produccion_linea_ids.filtered(
            lambda l: l.concepto == "ingreso" and l.fecha_mes == self.mes_actual)
        self.assertTrue(l_pasado.mes_cerrado)
        self.assertFalse(l_actual.mes_cerrado)
        with self.assertRaises(ValidationError):
            l_pasado.importe_previsto = 1.0

    # ------------------------------------------------------------ puerta de atrás
    def test_carga_inicial_requiere_grupo(self):
        p = self._proyecto()
        user = self.env["res.users"].create({
            "name": "PM sin permiso", "login": "pm_sin_permiso_test",
            "group_ids": [(6, 0, [self.env.ref("project.group_project_manager").id])],
        })
        self.assertFalse(user.has_group("aunna_produccion.group_produccion_carga_inicial"))
        # el modelo del asistente ni siquiera es accesible sin el grupo
        from odoo.exceptions import AccessError
        with self.assertRaises(AccessError):
            self.env["produccion.carga.inicial"].with_user(user).create({
                "project_id": p.id, "motivo": "x"})

    def test_carga_inicial_manual_registra_y_reparte(self):
        p = self._proyecto(tipo="closed", tcv=6000.0)
        # con costes previstos solo en meses abiertos, el reparto por avance de
        # coste concentra el ingreso ahí; los cerrados están a 0
        for l in p.produccion_linea_ids:
            if l.concepto == "horas_internas" and l.fecha_mes >= self.mes_actual:
                l.with_context(sin_recalculo=True, permitir_meses_pasados=PERMISO_MESES_CERRADOS).importe_previsto = -100.0
        p._recalcular_ingresos()
        ing = self._ingresos(p)
        self.assertAlmostEqual(sum(v for m, v in ing.items() if m < self.mes_actual), 0.0, 2)

        self.env.user.group_ids |= self.env.ref("aunna_produccion.group_produccion_carga_inicial")
        Wiz = self.env["produccion.carga.inicial"].with_context(default_project_id=p.id)
        w = Wiz.create({"project_id": p.id, "motivo": "Carga inicial de prueba"})
        # por defecto: solo INGRESO de meses cerrados → 3 celdas
        self.assertEqual(len(w.linea_ids), 3)
        self.assertTrue(all(c.mes_cerrado for c in w.linea_ids))
        for c in w.linea_ids:
            c.importe_nuevo = 500.0
        w.action_aplicar()

        ing = self._ingresos(p)
        for m, v in ing.items():
            if m < self.mes_actual:
                self.assertAlmostEqual(v, 500.0, 2)
        # el TCV restante para los abiertos es 6.000 − 1.500 = 4.500
        self.assertAlmostEqual(sum(v for m, v in ing.items() if m >= self.mes_actual), 4500.0, 2)
        # registro de auditoría: 3 filas con anterior 0 y nuevo 500
        ajustes = self.env["produccion.plan.ajuste"].search([("project_id", "=", p.id)])
        self.assertEqual(len(ajustes), 3)
        self.assertTrue(all(a.importe_anterior == 0.0 and a.importe_nuevo == 500.0 for a in ajustes))
        self.assertTrue(all(a.user_id == self.env.user for a in ajustes))
        # y nota en el chatter del proyecto
        self.assertTrue(p.message_ids.filtered(lambda m: "Carga inicial" in (m.body or "")))

    def test_carga_inicial_no_supera_tcv(self):
        p = self._proyecto(tipo="time_materials", tcv=1000.0)
        self.env.user.group_ids |= self.env.ref("aunna_produccion.group_produccion_carga_inicial")
        w = self.env["produccion.carga.inicial"].with_context(default_project_id=p.id).create(
            {"project_id": p.id, "motivo": "demasiado"})
        for c in w.linea_ids:
            c.importe_nuevo = 900.0   # 3 × 900 > 1.000
        self.assertTrue(w.aviso_tcv)
        with self.assertRaises(UserError):
            w.action_aplicar()

    def test_carga_inicial_excel(self):
        import base64
        import io
        import openpyxl
        p = self._proyecto(tipo="time_materials", tcv=9000.0)
        self.env.user.group_ids |= self.env.ref("aunna_produccion.group_produccion_carga_inicial")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=5, column=1, value="Concepto")
        meses = sorted({l.fecha_mes for l in p.produccion_linea_ids})
        for j, m in enumerate(meses):
            ws.cell(row=5, column=2 + j, value=m.strftime("%m/%Y"))
        ws.cell(row=6, column=1, value="INGRESO")
        for j, m in enumerate(meses):
            ws.cell(row=6, column=2 + j, value=1000.0)
        buf = io.BytesIO()
        wb.save(buf)
        w = self.env["produccion.carga.inicial"].with_context(default_project_id=p.id).create({
            "project_id": p.id, "motivo": "carga excel", "modo": "excel",
            "ambito": "todos", "fichero": base64.b64encode(buf.getvalue())})
        w.action_aplicar()
        ing = self._ingresos(p)
        self.assertEqual(len(ing), 6)
        for v in ing.values():
            self.assertAlmostEqual(v, 1000.0, 2)
        ajustes = self.env["produccion.plan.ajuste"].search([("project_id", "=", p.id)])
        self.assertEqual(len(ajustes), 6)
        self.assertTrue(all(a.origen == "excel" for a in ajustes))


    # ------------------------------------------------------------ hallazgos de la revisión 18/08
    def test_recurrente_con_carga_en_cerrados_no_supera_tcv(self):
        """Recurrente: la cuota se reduce si lo cargado en meses cerrados no deja
        sitio; el total nunca supera el TCV y no queda «sin distribuir»."""
        p = self._proyecto(tipo="recurring", tcv=6000.0)
        self.env.user.group_ids |= self.env.ref("aunna_produccion.group_produccion_carga_inicial")
        w = self.env["produccion.carga.inicial"].with_context(default_project_id=p.id).create(
            {"project_id": p.id, "motivo": "historico"})
        for c in w.linea_ids:
            c.importe_nuevo = 1500.0   # 3 x 1.500 = 4.500 en cerrados
        self.assertFalse(w.aviso_tcv)
        w.action_aplicar()
        ing = self._ingresos(p)
        self.assertAlmostEqual(sum(ing.values()), 6000.0, 2)   # 4.500 + 3 x 500
        for m, v in ing.items():
            self.assertAlmostEqual(v, 1500.0 if m < self.mes_actual else 500.0, 2)
        self.assertAlmostEqual(p.produccion_tcv_restante, 0.0, 2)

    def test_recurrente_sin_carga_mantiene_cuota(self):
        """Sin nada en cerrados, la cuota sigue siendo TCV/n (comportamiento previo)."""
        p = self._proyecto(tipo="recurring", tcv=6000.0)
        ing = self._ingresos(p)
        for m, v in ing.items():
            self.assertAlmostEqual(v, 0.0 if m < self.mes_actual else 1000.0, 2)

    def test_aviso_asistente_no_es_falso_en_cerrado(self):
        """I3: en proyecto cerrado los abiertos se rehacen al aplicar → sin aviso."""
        p = self._proyecto(tipo="closed", tcv=6000.0)
        for l in p.produccion_linea_ids:
            if l.concepto == "horas_internas" and l.fecha_mes >= self.mes_actual:
                l.with_context(sin_recalculo=True, permitir_meses_pasados=PERMISO_MESES_CERRADOS).importe_previsto = -100.0
        p._recalcular_ingresos()
        self.env.user.group_ids |= self.env.ref("aunna_produccion.group_produccion_carga_inicial")
        w = self.env["produccion.carga.inicial"].with_context(default_project_id=p.id).create(
            {"project_id": p.id, "motivo": "x"})
        for c in w.linea_ids:
            c.importe_nuevo = 500.0
        self.assertFalse(w.aviso_tcv, "los abiertos (6.000) se rehacen: no debe avisar")
        for c in w.linea_ids:
            c.importe_nuevo = 3000.0   # 9.000 en cerrados > TCV: ahora si
        self.assertTrue(w.aviso_tcv)

    def test_asistente_no_ofrece_ingreso_abierto_en_automaticos(self):
        p = self._proyecto(tipo="recurring", tcv=6000.0)
        self.env.user.group_ids |= self.env.ref("aunna_produccion.group_produccion_carga_inicial")
        w = self.env["produccion.carga.inicial"].with_context(default_project_id=p.id).create(
            {"project_id": p.id, "motivo": "x", "ambito": "todos"})
        self.assertTrue(all(c.mes_cerrado for c in w.linea_ids))

    def test_conceptos_vigentes_incluyen_selection_add(self):
        Linea = self.env["produccion.plan.linea"]
        claves = Linea._claves_coste_vigentes()
        for c in ("horas_internas", "materiales", "otros"):
            self.assertIn(c, claves)
        self.assertNotIn("ingreso", claves)
        if "licencias_software" in dict(Linea._fields["concepto"].selection):
            self.assertIn("licencias_software", claves)
            p = self._proyecto(tipo="time_materials")
            self.assertTrue(p.produccion_linea_ids.filtered(
                lambda l: l.concepto == "licencias_software"))

    def test_contexto_por_rpc_no_salta_meses_cerrados(self):
        """Un usuario normal no puede saltarse la regla 3 pasando el contexto
        `permitir_meses_pasados` por la API: solo vale el testigo interno del
        servidor. Y la carga inicial autorizada sí puede."""
        p = self._proyecto(tipo="time_materials")
        pm = self.env["res.users"].create({
            "name": "PM RPC", "login": "pm_rpc_test",
            "group_ids": [(6, 0, [self.env.ref("project.group_project_manager").id])],
        })
        cerrada = p.produccion_linea_ids.filtered(
            lambda l: l.concepto == "ingreso" and l.fecha_mes == self.ini)
        self.assertEqual(len(cerrada), 1)
        self.assertTrue(cerrada.mes_cerrado)
        # lo que un cliente RPC puede mandar: True, 1, el NOMBRE del testigo...
        for falso in (True, 1, "PERMISO_MESES_CERRADOS", "object()"):
            with self.assertRaises(ValidationError):
                cerrada.with_user(pm).with_context(permitir_meses_pasados=falso).write(
                    {"importe_previsto": 123.0})
        # con el grupo, por el asistente, sí
        pm.group_ids |= self.env.ref("aunna_produccion.group_produccion_carga_inicial")
        w = self.env["produccion.carga.inicial"].with_user(pm).with_context(
            default_project_id=p.id).create({"project_id": p.id, "motivo": "ok"})
        for c in w.linea_ids:
            c.importe_nuevo = 100.0
        w.action_aplicar()
        self.assertAlmostEqual(cerrada.importe_previsto, 100.0, 2)
