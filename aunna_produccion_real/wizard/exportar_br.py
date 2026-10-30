# -*- coding: utf-8 -*-
"""Exportación Excel con el formato visual del BR corporativo: pasado cian,
mes en curso azul intenso, futuro morado, agregados Q/H/FY morado oscuro."""
import base64
import io

from odoo import fields, models
from odoo.exceptions import UserError

# Paleta BR con el mes abierto diferenciado, igual que en pantalla.
COLOR_PASADO = "6DD9FF"
COLOR_ACTUAL = "2563EB"
COLOR_FUTURO = "BD92DE"
COLOR_AGREGADO = "7030A0"
COLOR_GRUPOS = {
    "avance": "1F4E79",
    "costes": "843C39",
    "pm": "375623",
    "wip": "BF6000",
}


class ProduccionExportarBR(models.TransientModel):
    _name = "produccion.exportar.br"
    _description = "Exportar Seguimiento económico a Excel"

    project_id = fields.Many2one("project.project", required=True)
    anio = fields.Integer(
        string="Año (vacío = todo el periodo)",
        help="Deje vacío para exportar el periodo completo del proyecto.",
    )
    incluir_agregados = fields.Boolean(string="Incluir Q / H / FY", default=True)
    fichero = fields.Binary(readonly=True)
    nombre_fichero = fields.Char(readonly=True)

    def action_exportar(self):
        self.ensure_one()
        try:
            import openpyxl
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        except ImportError:
            raise UserError("Falta la librería openpyxl en el servidor.")

        datos = self.project_id.get_produccion_br_data(
            anio=self.anio or None
        )
        if not datos["meses"]:
            raise UserError(
                "El proyecto no tiene planificación. Genérela primero en la "
                "pestaña Producción."
            )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Seguimiento economico"
        borde = Border(*(Side(style="thin", color="D8DEE6"),) * 4)
        negrita = Font(bold=True)
        blanco = Font(bold=True, color="FFFFFF")
        centro = Alignment(horizontal="center")
        derecha = Alignment(horizontal="right")

        def fill(color):
            return PatternFill("solid", fgColor=color)

        # ---- columnas: meses + agregados (misma lógica que la vista) ----
        MESES_ES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun",
                    "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
        columnas = []  # (etiqueta, tipo, meses_que_agrega)
        for clave in datos["meses"]:
            mm = int(clave[5:7])
            etiqueta = f"{MESES_ES[mm-1]} {clave[2:4]}"
            if datos["estado_mes"][clave] == "actual":
                etiqueta += " · En curso"
            columnas.append((etiqueta, clave, [clave]))
            if self.incluir_agregados and mm in (3, 6, 9, 12):
                anio_m = clave[:4]
                del_anio = [m for m in datos["meses"] if m.startswith(anio_m)]
                q = [m for m in del_anio if mm - 3 < int(m[5:7]) <= mm]
                columnas.append((f"Q{mm // 3}", "agregado", q))
                if mm == 6:
                    columnas.append(
                        ("H1", "agregado",
                         [m for m in del_anio if int(m[5:7]) <= 6])
                    )
                if mm == 12:
                    columnas.append(
                        ("H2", "agregado",
                         [m for m in del_anio if int(m[5:7]) > 6])
                    )
                    columnas.append(("FY", "agregado", del_anio))
        if self.incluir_agregados:
            columnas.append(("Total", "agregado", datos["meses"]))

        # ---- cabecera ----
        ws.cell(row=1, column=1, value=self.project_id.name).font = Font(
            bold=True, size=14
        )
        ws.cell(row=2, column=1, value="TCV")
        c = ws.cell(row=2, column=2, value=datos["tcv"])
        c.number_format = "#,##0.00"
        ws.cell(
            row=2,
            column=3,
            value="(fuente: %s)"
            % ("ficha del proyecto" if datos["tcv_fuente"] == "ficha" else "manual"),
        )
        ws.cell(row=3, column=1, value="Último recálculo guardado")
        ws.cell(row=3, column=2, value=datos["last_update"] or "sin recalcular")
        if datos["costes_actuales_hasta"]:
            ws.cell(row=4, column=1, value=(
                "Mes en curso abierto: costes reales hasta %s; ingreso previsto."
                % datos["costes_actuales_hasta"]
            ))

        fila = 5
        ws.cell(row=fila, column=1, value="Concepto").font = negrita
        for j, (etiqueta, tipo, _m) in enumerate(columnas, start=2):
            celda = ws.cell(row=fila, column=j, value=etiqueta)
            celda.font = negrita if tipo != "agregado" else blanco
            celda.alignment = centro
            celda.border = borde
            if tipo == "agregado":
                celda.fill = fill(COLOR_AGREGADO)
            else:
                estado = datos["estado_mes"][tipo]
                if estado == "actual":
                    celda.font = blanco
                celda.fill = fill(
                    COLOR_PASADO if estado == "pasado"
                    else COLOR_ACTUAL if estado == "actual"
                    else COLOR_FUTURO
                )
        fila += 1

        def valor_mes(tipo_fila, concepto, mes):
            if tipo_fila == "avance":
                return datos["avance"].get(mes, {}).get("valor", 0.0)
            if tipo_fila == "coste":
                for f in datos["costes"]:
                    if f["concepto"] == concepto:
                        return f["celdas"].get(mes, {}).get("valor", 0.0)
                return 0.0
            if tipo_fila == "pm":
                return datos["pm"].get(mes, 0.0)
            if tipo_fila == "fact":
                return datos["facturacion"].get(mes, 0.0)
            if tipo_fila == "wip":
                return datos["wip"].get(mes, 0.0)
            return 0.0

        def escribir_fila(etiqueta, tipo_fila, concepto=None, pct=False, negrita_fila=False):
            nonlocal fila
            c0 = ws.cell(row=fila, column=1, value=etiqueta)
            c0.border = borde
            if negrita_fila:
                c0.font = negrita
            for j, (_e, tipo, meses_col) in enumerate(columnas, start=2):
                if tipo_fila == "wip":
                    v = valor_mes("wip", None, meses_col[-1]) if meses_col else 0.0
                elif pct:
                    if tipo != "agregado":
                        v = (
                            datos["pct_pm_acum" if "acum" in etiqueta.lower() else "pct_pm"]
                            .get(meses_col[0], 0.0)
                        )
                    elif "acum" in etiqueta.lower():
                        v = datos["pct_pm_acum"].get(meses_col[-1], 0.0) if meses_col else 0.0
                    else:
                        av = sum(valor_mes("avance", None, m) for m in meses_col)
                        pmv = sum(valor_mes("pm", None, m) for m in meses_col)
                        v = (pmv / av * 100.0) if av else 0.0
                else:
                    v = sum(valor_mes(tipo_fila, concepto, m) for m in meses_col)
                celda = ws.cell(row=fila, column=j, value=(v / 100.0) if pct else v)
                celda.number_format = "0.00%" if pct else "#,##0.00"
                celda.alignment = derecha
                celda.border = borde
                if negrita_fila or tipo == "agregado":
                    celda.font = negrita
                if tipo == "agregado":
                    celda.fill = fill("ECE2F4")
                else:
                    estado = datos["estado_mes"][meses_col[0]]
                    celda.fill = fill(
                        "EEFAFF" if estado == "pasado"
                        else "DBEAFE" if estado == "actual"
                        else "F7F0FC"
                    )
            fila += 1

        def cabecera_grupo(texto, clave_color):
            nonlocal fila
            celda = ws.cell(row=fila, column=1, value=texto)
            celda.font = blanco
            celda.fill = fill(COLOR_GRUPOS[clave_color])
            for j in range(2, len(columnas) + 2):
                ws.cell(row=fila, column=j).fill = fill(COLOR_GRUPOS[clave_color])
            fila += 1

        cabecera_grupo("AVANCE — ingreso reconocido (€)", "avance")
        escribir_fila("Avance", "avance", negrita_fila=True)
        cabecera_grupo("COSTES (€)", "costes")
        for f in datos["costes"]:
            escribir_fila(f["etiqueta"], "coste", f["concepto"])
        cabecera_grupo("PM (RENTABILIDAD)", "pm")
        escribir_fila("PM", "pm", negrita_fila=True)
        escribir_fila("%PM", None, pct=True)
        escribir_fila("%PM (acumulado)", None, pct=True)
        cabecera_grupo("FACTURACIÓN Y WIP", "wip")
        escribir_fila("Facturación", "fact")
        escribir_fila("WIP (avance − facturado)", "wip", negrita_fila=True)

        ws.column_dimensions["A"].width = 30
        for j in range(2, len(columnas) + 2):
            ws.column_dimensions[
                openpyxl.utils.get_column_letter(j)
            ].width = 24 if datos["estado_mes"].get(columnas[j-2][1]) == "actual" else 12
        ws.freeze_panes = "B6"

        # ============ Hoja 2: Previsto vs Real (solo meses, sin agregados) ====
        ws2 = wb.create_sheet("Previsto vs Real")
        ws2.cell(row=1, column=1,
                 value="Previsto vs Real por concepto y mes (€)").font = Font(
            bold=True, size=12)
        ws2.cell(row=2, column=1, value=(
            "Avance: la columna Real solo muestra meses pasados; "
            "el ingreso del mes en curso sigue siendo previsión."
        ))
        fila = 3
        ws2.cell(row=fila, column=1, value="Concepto").font = negrita
        col = 2
        for clave in datos["meses"]:
            mm = int(clave[5:7])
            estado = datos["estado_mes"][clave]
            color = (COLOR_PASADO if estado == "pasado"
                     else COLOR_ACTUAL if estado == "actual" else COLOR_FUTURO)
            etiqueta = f"{MESES_ES[mm-1]} {clave[2:4]}"
            if estado == "actual":
                etiqueta += " · En curso"
            c = ws2.cell(row=fila, column=col, value=etiqueta)
            c.font = blanco if estado == "actual" else negrita
            c.alignment = centro
            c.fill = fill(color)
            ws2.merge_cells(start_row=fila, start_column=col,
                            end_row=fila, end_column=col + 2)
            for k, sub in enumerate(["Prev.", "Real", "Δ"]):
                sc = ws2.cell(row=fila + 1, column=col + k, value=sub)
                sc.font = blanco if estado == "actual" else negrita
                sc.alignment = centro
                sc.fill = fill(color)
            col += 3

        def fila_pvr(etiqueta, prev_por_mes, real_por_mes, negrita_fila=False):
            nonlocal fila_datos
            c0 = ws2.cell(row=fila_datos, column=1, value=etiqueta)
            c0.border = borde
            if negrita_fila:
                c0.font = negrita
            col = 2
            for clave in datos["meses"]:
                prev = prev_por_mes.get(clave, 0.0)
                real = real_por_mes.get(clave, 0.0)
                for k, v in enumerate([prev, real, real - prev if real is not None else None]):
                    cd = ws2.cell(row=fila_datos, column=col + k, value=v)
                    cd.number_format = "#,##0.00"
                    cd.alignment = derecha
                    cd.border = borde
                    if negrita_fila:
                        cd.font = negrita
                col += 3
            fila_datos += 1

        fila_datos = fila + 2
        fila_pvr(
            "Avance (ingreso)",
            {m: datos["avance"][m]["previsto"] for m in datos["meses"]},
            {m: datos["avance"][m]["valor"]
             if datos["estado_mes"][m] == "pasado" else None
             for m in datos["meses"]},
            negrita_fila=True,
        )
        for f in datos["costes"]:
            fila_pvr(
                f["etiqueta"],
                {m: f["celdas"][m]["previsto"] for m in datos["meses"]},
                {m: f["celdas"][m]["real"] for m in datos["meses"]},
            )
        ws2.column_dimensions["A"].width = 30
        ws2.freeze_panes = "B5"

        # ============ Hoja 3: detalle de horas =================================
        motor = self.env["produccion.real.motor"]
        d_ini = fields.Date.to_date(datos["meses"][0] + "-01")
        import calendar as _cal

        ult = datos["meses"][-1]
        d_fin = fields.Date.to_date(
            "%s-%02d" % (ult, _cal.monthrange(int(ult[:4]), int(ult[5:7]))[1])
        )
        horas = motor.detalle_horas(self.project_id, d_ini, d_fin)
        ws3 = wb.create_sheet("Horas (detalle)")
        ws3.cell(row=1, column=1, value="Partes de horas del periodo — solo las "
                 "VALIDADAS computan como coste").font = Font(bold=True, size=12)
        cab = ["Fecha", "Empleado", "Descripción", "Horas", "Coste €",
               "Tipo", "Validada", "Computa"]
        for j, t in enumerate(cab, start=1):
            c = ws3.cell(row=3, column=j, value=t)
            c.font = blanco
            c.fill = fill(COLOR_GRUPOS["costes"])
        for i, h in enumerate(horas, start=4):
            ws3.cell(row=i, column=1, value=h["fecha"].strftime("%d/%m/%Y"))
            ws3.cell(row=i, column=2, value=h["empleado"])
            ws3.cell(row=i, column=3, value=h["descripcion"][:80])
            ws3.cell(row=i, column=4, value=h["horas"]).number_format = "0.0"
            ws3.cell(row=i, column=5, value=h["coste"]).number_format = "#,##0.00"
            ws3.cell(row=i, column=6, value=h["tipo"])
            ws3.cell(row=i, column=7, value="Sí" if h["validada"] else "No")
            cc = ws3.cell(row=i, column=8, value="Sí" if h["computa"] else "No")
            if not h["computa"]:
                for j in range(1, 9):
                    ws3.cell(row=i, column=j).font = Font(color="999999")
        tot_c = sum(h["coste"] for h in horas if h["computa"])
        tot_h = sum(h["horas"] for h in horas if h["computa"])
        frow = 4 + len(horas) + 1
        ws3.cell(row=frow, column=3, value="TOTAL computable").font = negrita
        ws3.cell(row=frow, column=4, value=tot_h).font = negrita
        c = ws3.cell(row=frow, column=5, value=tot_c)
        c.font = negrita
        c.number_format = "#,##0.00"
        for letra, ancho in zip("ABCDEFGH", (11, 22, 50, 8, 11, 15, 9, 9)):
            ws3.column_dimensions[letra].width = ancho
        ws3.freeze_panes = "A4"

        # ============ Hoja 4: detalle de compras ===============================
        compras = motor.detalle_compras(self.project_id, d_ini, d_fin)
        ws4 = wb.create_sheet("Otros costes (detalle)")
        ws4.cell(row=1, column=1, value="Costes del periodo distintos de horas: "
                 "recepciones/devoluciones de compra, facturas de proveedor, "
                 "gastos de empleado y entregas de almacén (en gris: no computa, "
                 "ya contado por su recepción)").font = Font(bold=True, size=12)
        cab = ["Fecha", "Origen", "Documento", "Tercero", "Referencia",
               "Descripción", "Tipo de coste", "Importe €", "Computa"]
        for j, t in enumerate(cab, start=1):
            c = ws4.cell(row=3, column=j, value=t)
            c.font = blanco
            c.fill = fill(COLOR_GRUPOS["costes"])
        for i, x in enumerate(compras, start=4):
            ws4.cell(row=i, column=1, value=x["fecha"].strftime("%d/%m/%Y"))
            ws4.cell(row=i, column=2, value=x["origen"])
            ws4.cell(row=i, column=3, value=x["documento"])
            ws4.cell(row=i, column=4, value=x["tercero"])
            ws4.cell(row=i, column=5, value=x["referencia"])
            ws4.cell(row=i, column=6, value=x["descripcion"][:60])
            ws4.cell(row=i, column=7, value=x["tipo"])
            ws4.cell(row=i, column=8, value=x["importe"]).number_format = "#,##0.00"
            computa = x.get("computa", True)
            ws4.cell(row=i, column=9, value="Sí" if computa else "No")
            if not computa:
                for j in range(1, 10):
                    ws4.cell(row=i, column=j).font = Font(color="999999")
        if compras:
            frow = 4 + len(compras) + 1
            ws4.cell(row=frow, column=7, value="TOTAL computable").font = negrita
            c = ws4.cell(row=frow, column=8,
                         value=sum(x["importe"] for x in compras if x.get("computa", True)))
            c.font = negrita
            c.number_format = "#,##0.00"
        for letra, ancho in zip("ABCDEFGHI", (11, 26, 16, 24, 18, 40, 22, 12, 9)):
            ws4.column_dimensions[letra].width = ancho
        ws4.freeze_panes = "A4"

        buf = io.BytesIO()
        wb.save(buf)
        nombre = "BR_%s_%s.xlsx" % (
            (self.project_id.name or "proyecto").replace("/", "-")[:40],
            fields.Date.context_today(self).strftime("%Y%m%d"),
        )
        self.write(
            {
                "fichero": base64.b64encode(buf.getvalue()),
                "nombre_fichero": nombre,
            }
        )
        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "res_id": self.id,
            "view_mode": "form",
            "target": "new",
        }
